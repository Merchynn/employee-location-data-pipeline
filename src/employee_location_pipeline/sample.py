from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook

HEADERS = [
    "employee_id", "employee_code", "full_name", "branch", "department",
    "cost_center", "admission_date", "postal_code", "street", "street_number",
    "address_complement", "neighborhood", "city", "state", "declared_commute_time",
    "latitude", "longitude", "route_distance_km", "route_duration_text",
    "uses_train", "uses_walking", "uses_subway", "uses_bus", "uses_car",
    "source_updated_at",
]

ROWS = [
    ["E001", "A1001", "Ana Example", "HEADQUARTERS", "Analytics", "CC100", "2024-02-01", "01310100", "Avenida Example", "100", "Apto 10", "Centro", "São Paulo", "SP", "1:05", -23.5614, -46.6560, 8.2, "35 minutos", 0, 1, 1, 0, 0, "2026-07-01"],
    ["E002", "A1002", "Bruno Sample", "HEADQUARTERS", "Finance", "CC200", "2023-08-14", "04001000", "Rua Demonstracao", "250", "", "Vila Nova", "São Paulo", "SP", "55 min", -23.5780, -46.6430, 6.4, "28 minutos", 0, 1, 0, 1, 0, "2026-07-01"],
    ["E003", "A1003", "Carla Demo", "HEADQUARTERS", "People", "CC300", "2022-11-03", "05001000", "Alameda Synthetic", "75", "Casa 2", "Oeste", "São Paulo", "SP", "1h 30 min", -23.5350, -46.6900, 12.8, "50 minutos", 1, 0, 1, 1, 0, "2026-07-01"],
    ["E004", "A1004", "Diego Test", "HEADQUARTERS", "Operations", "CC400", "2025-01-20", "03001000", "Rua Pipeline", "900", "", "Leste", "São Paulo", "SP", "45", -23.5450, -46.6100, 5.9, "25 minutos", 0, 1, 0, 1, 0, "2026-07-01"],
    ["E005", "B2001", "Elisa Mock", "REGIONAL_OFFICE", "Analytics", "CC110", "2021-06-18", "78005000", "Avenida Sample", "45", "", "Centro", "Cuiabá", "MT", "0:35", -15.5980, -56.0920, 2.1, "12 minutos", 0, 1, 0, 0, 1, "2026-07-01"],
    ["E006", "B2002", "Fabio Example", "REGIONAL_OFFICE", "Sales", "CC500", "2024-09-09", "78010000", "Rua Metrics", "123", "Bloco B", "Norte", "Cuiabá", "MT", "50 min", -15.6200, -56.1100, 4.7, "20 minutos", 0, 1, 0, 1, 0, "2026-07-01"],
    ["E007", "B2003", "Gabriela Sample", "REGIONAL_OFFICE", "Engineering", "CC600", "2020-04-27", "78020000", "Travessa Data", "10", "", "Sul", "Cuiabá", "MT", "1h 10 min", -15.6400, -56.0800, 7.9, "32 minutos", 0, 0, 0, 1, 1, "2026-07-01"],
    ["E008", "B2004", "Henrique Demo", "REGIONAL_OFFICE", "Legal", "CC700", "2025-03-11", "78030000", "Alameda Quality", "500", "Apto 5", "Leste", "Cuiabá", "MT", "40 min", -15.5850, -56.1200, 5.1, "22 minutos", 0, 1, 0, 1, 0, "2026-07-01"],
]


def ensure_sample_excel(path: Path) -> Path:
    if path.exists():
        return path
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Employees"
    worksheet.append(HEADERS)
    for row in ROWS:
        worksheet.append(row)
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = f"A1:Y{worksheet.max_row}"
    instructions = workbook.create_sheet("Instructions")
    instructions.append(["Synthetic employee data generated for the public demo."])
    instructions.append(["Do not replace this file with real employee data in a public repository."])
    branches = workbook.create_sheet("Branches")
    branches.append(["branch", "latitude", "longitude"])
    branches.append(["HEADQUARTERS", -23.5505, -46.6333])
    branches.append(["REGIONAL_OFFICE", -15.6014, -56.0979])
    workbook.save(path)
    return path
